"""Fetch the latest Claude weekly tarot crawl result and normalize it for the dashboard.

The Claude cloud routine "주간 타로 크롤링 (월요일 아침)" uploads one Google Sheet per week into the
Drive folder "주간 타로 크롤링". Each sheet is named "레타 주간 타로 크롤링_YYYY-MM-DD" and has four tabs:

    raw_crawl, topic_classification, content_ideas, summary

This script reads the newest sheet (or a local xlsx copy) and writes ``data/claude_weekly_tarot.json``
in the schema that ``scripts/generate_site.py`` renders.

Usage:
    python scripts/fetch_claude_crawl.py                 # newest sheet in the Drive folder
    python scripts/fetch_claude_crawl.py --spreadsheet-id <id>
    python scripts/fetch_claude_crawl.py --xlsx path/to/레타_주간_타로_크롤링_2026-10-05.xlsx

Environment:
    GOOGLE_SERVICE_ACCOUNT_FILE   service account JSON (same one the pipeline uses)
    CLAUDE_CRAWL_DRIVE_FOLDER_ID  Drive folder that holds the weekly sheets
    CLAUDE_CRAWL_SPREADSHEET_ID   optional: read this sheet instead of searching the folder
    CLAUDE_CRAWL_TITLE_PREFIX     optional: sheet title prefix (default "레타 주간 타로 크롤링")
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from dotenv import load_dotenv


ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = ROOT / "data" / "claude_weekly_tarot.json"
KST = timezone(timedelta(hours=9))

DEFAULT_FOLDER_ID = "1_5FYHF5haixRGH8Dh-jolrwWf3jivOjz"
DEFAULT_TITLE_PREFIX = "레타 주간 타로 크롤링"
TABS = ("raw_crawl", "topic_classification", "content_ideas", "summary")
OPTIONAL_TABS = ("reading_analysis",)
TRANSCRIPT_CACHE_PATH = ROOT / "data" / "transcript_analysis.json"
VIDEO_ID_RE = re.compile(r"(?:v=|youtu\.be/|/shorts/|/embed/)([\w-]{11})")
SCOPES = [
    "https://www.googleapis.com/auth/drive.readonly",
    "https://www.googleapis.com/auth/spreadsheets.readonly",
]
DATE_IN_TITLE = re.compile(r"(\d{4}-\d{2}-\d{2})")


def log(message: str) -> None:
    print(f"[fetch_claude_crawl] {message}", file=sys.stderr)


# --------------------------------------------------------------------------- helpers


def rows_to_records(values: List[List[Any]]) -> List[Dict[str, Any]]:
    """Turn a header row plus data rows into dicts, skipping fully blank rows."""
    if not values:
        return []
    headers = [str(cell).strip() for cell in values[0]]
    records: List[Dict[str, Any]] = []
    for row in values[1:]:
        if not any(str(cell).strip() for cell in row):
            continue
        record: Dict[str, Any] = {}
        for index, header in enumerate(headers):
            if not header:
                continue
            record[header] = row[index] if index < len(row) else ""
        records.append(record)
    return records


def to_int(value: Any) -> Optional[int]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    digits = re.sub(r"[^\d]", "", str(value))
    return int(digits) if digits else None


def to_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value))
    return float(match.group()) if match else None


def split_urls(value: Any) -> List[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value or "")
    return [part.strip() for part in re.split(r"[\s,;\n]+", text) if part.strip().startswith("http")]


def clean_text(value: Any) -> str:
    return str(value if value is not None else "").strip()


def normalize_raw(records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows = []
    for record in records:
        rows.append(
            {
                "run_date": clean_text(record.get("run_date")),
                "platform": clean_text(record.get("platform")) or "youtube",
                "keyword": clean_text(record.get("keyword")),
                "title": clean_text(record.get("title")),
                "author": clean_text(record.get("author")),
                "published_at": clean_text(record.get("published_at")),
                "url": clean_text(record.get("url")),
                "views": to_int(record.get("views")),
                "likes": to_int(record.get("likes")),
                "duration_min": to_float(record.get("duration_min")),
                "exclude_reason": clean_text(record.get("exclude_reason")),
                "source_type": clean_text(record.get("source_type")),
            }
        )
    return rows


def normalize_classification(records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows = []
    for record in records:
        row = {
            "major_topic": clean_text(record.get("major_topic")),
            "sub_topic": clean_text(record.get("sub_topic")),
            "content_angle": clean_text(record.get("content_angle")),
            "target_emotion": clean_text(record.get("target_emotion")),
            "confidence": to_float(record.get("confidence")) if to_float(record.get("confidence")) is not None else clean_text(record.get("confidence")),
        }
        # Some weeks carry extra columns (title/url/keyword). Keep them if present.
        for optional in ("title", "url", "keyword"):
            if record.get(optional):
                row[optional] = clean_text(record.get(optional))
        rows.append(row)
    return rows


def normalize_ideas(records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows = []
    for record in records:
        rows.append(
            {
                "idea_title": clean_text(record.get("idea_title")),
                "major_topic": clean_text(record.get("major_topic")),
                "sub_topic": clean_text(record.get("sub_topic")),
                "hook": clean_text(record.get("hook")),
                "summary": clean_text(record.get("summary")),
                "recommended_format": clean_text(record.get("recommended_format")),
                "source_urls": split_urls(record.get("source_urls")),
                "priority_score": to_float(record.get("priority_score")),
            }
        )
    rows.sort(key=lambda row: row["priority_score"] or 0, reverse=True)
    return rows


def normalize_summary(values: List[List[Any]]) -> Dict[str, Any]:
    """The summary tab is free-form: label/value pairs with section headings in between.

    We keep every non-empty row as a (label, value) pair and additionally pick out the
    rows that look like the "3대 발견" findings and the top recommendation.
    """
    pairs: List[Dict[str, str]] = []
    for row in values:
        cells = [clean_text(cell) for cell in row]
        cells = [cell for cell in cells if cell]
        if not cells:
            continue
        label = cells[0]
        value = " ".join(cells[1:]) if len(cells) > 1 else ""
        pairs.append({"label": label, "value": value})

    def as_finding(label: str, value: str) -> str:
        title = re.sub(r"^(\d+[.)]\s*|[①②③④⑤]\s*|발견\s*\d\s*[.:)]?\s*)", "", label).strip()
        if title and value:
            return f"{title} — {value}"
        return title or value

    findings: List[str] = []
    recommendation = ""
    in_findings = False
    for pair in pairs:
        label, value = pair["label"], pair["value"]
        text = f"{label} {value}".strip()
        if "발견" in label and not value:
            in_findings = True
            continue
        if in_findings:
            if re.match(r"^(\d+[.)]|[①②③④⑤]|발견\s*\d)", label) or (value and len(findings) < 3):
                findings.append(as_finding(label, value))
                continue
            in_findings = False
        if re.match(r"^(발견|3대 발견)\s*\d", label) or re.match(r"^[①②③]", label):
            findings.append(as_finding(label, value))
            continue
        if ("1순위" in label or "추천" in label) and not recommendation:
            recommendation = value or text

    return {"rows": pairs, "findings": findings[:3], "top_recommendation": recommendation}


def split_list(value: Any, sep: str) -> List[str]:
    return [part.strip() for part in str(value or "").split(sep) if part.strip()]


def normalize_reading_analysis(records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Rows the Claude routine writes to the optional `reading_analysis` tab (one per analyzed video)."""
    rows = []
    for record in records:
        url = clean_text(record.get("url"))
        match = VIDEO_ID_RE.search(url)
        video_id = clean_text(record.get("video_id")) or (match.group(1) if match else "")
        if not video_id:
            continue
        rows.append(
            {
                "video_id": video_id,
                "url": url or f"https://www.youtube.com/watch?v={video_id}",
                "title": clean_text(record.get("title")),
                "keyword": clean_text(record.get("keyword")),
                "transcript_chars": to_int(record.get("transcript_chars")),
                "analysis": {
                    "opening_hook": clean_text(record.get("opening_hook")),
                    "card_selection": {
                        "method": clean_text(record.get("card_method")) or "other",
                        "card_count": to_int(record.get("card_count")),
                        "staging": clean_text(record.get("card_staging")),
                    },
                    "reading_structure": split_list(record.get("reading_structure"), "→"),
                    "tone": clean_text(record.get("tone")),
                    "tone_tags": split_list(record.get("tone_tags"), ",")[:3],
                    "recurring_phrases": split_list(record.get("recurring_phrases"), ",")[:3],
                    "audience_engagement": clean_text(record.get("audience_engagement")),
                    "cta": clean_text(record.get("cta")),
                    "letter_concept_takeaway": clean_text(record.get("letter_concept_takeaway")),
                },
            }
        )
    return rows


def merge_reading_analysis_into_cache(rows: List[Dict[str, Any]], run_date: str) -> int:
    """Merge routine-made analyses into data/transcript_analysis.json unless a video is already analyzed."""
    if not rows:
        return 0
    cache: Dict[str, Any] = {"version": 1, "updated_at": "", "videos": {}}
    if TRANSCRIPT_CACHE_PATH.exists():
        try:
            loaded = load_json(TRANSCRIPT_CACHE_PATH)
            if isinstance(loaded, dict) and isinstance(loaded.get("videos"), dict):
                cache = loaded
        except Exception as error:  # noqa: BLE001
            log(f"transcript cache unreadable, recreating: {error}")
    merged = 0
    now = datetime.now(tz=KST).isoformat(timespec="seconds")
    for row in rows:
        entry = cache["videos"].get(row["video_id"]) or {}
        if entry.get("status") == "ok":
            continue
        entry.update(
            {
                "video_id": row["video_id"],
                "url": row["url"],
                "title": row["title"] or entry.get("title"),
                "keyword": row["keyword"] or entry.get("keyword"),
                "source": entry.get("source") or "claude",
                "status": "ok",
                "error": None,
                "analysis": row["analysis"],
                "model": f"claude-routine ({run_date})",
                "analyzed_at": now,
                "transcript": {**(entry.get("transcript") or {}), "chars": row["transcript_chars"] or (entry.get("transcript") or {}).get("chars", 0)},
            }
        )
        entry["transcript"].pop("text", None)
        cache["videos"][row["video_id"]] = entry
        merged += 1
    if merged:
        cache["updated_at"] = now
        TRANSCRIPT_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
        TRANSCRIPT_CACHE_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    return merged


def keyword_stats(raw_rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    buckets: Dict[str, Dict[str, Any]] = defaultdict(lambda: {"keyword": "", "collected": 0, "included": 0, "excluded": 0, "views": 0})
    for row in raw_rows:
        keyword = row["keyword"] or "미지정"
        bucket = buckets[keyword]
        bucket["keyword"] = keyword
        bucket["collected"] += 1
        if row["exclude_reason"]:
            bucket["excluded"] += 1
        else:
            bucket["included"] += 1
            bucket["views"] += row["views"] or 0
    return sorted(buckets.values(), key=lambda item: (item["included"], item["views"]), reverse=True)


def build_payload(
    tabs: Dict[str, List[List[Any]]],
    source: Dict[str, Any],
    archive: List[Dict[str, Any]],
) -> Dict[str, Any]:
    raw_rows = normalize_raw(rows_to_records(tabs.get("raw_crawl", [])))
    classification = normalize_classification(rows_to_records(tabs.get("topic_classification", [])))
    ideas = normalize_ideas(rows_to_records(tabs.get("content_ideas", [])))
    summary = normalize_summary(tabs.get("summary", []))
    reading_analysis = normalize_reading_analysis(rows_to_records(tabs.get("reading_analysis", [])))

    run_date = source.get("run_date") or next((row["run_date"] for row in raw_rows if row["run_date"]), "")
    source_types = sorted({row["source_type"] for row in raw_rows if row["source_type"]})
    has_metrics = any(row["views"] for row in raw_rows)

    return {
        "schema_version": 2,
        "generated_at": datetime.now(tz=KST).isoformat(timespec="seconds"),
        "run_date": run_date,
        "source": source,
        "collection": {
            "source_types": source_types,
            "has_view_metrics": has_metrics,
        },
        "stats": {
            "collected": len(raw_rows),
            "excluded": sum(1 for row in raw_rows if row["exclude_reason"]),
            "included": sum(1 for row in raw_rows if not row["exclude_reason"]),
            "ideas": len(ideas),
        },
        "summary": summary,
        "content_ideas": ideas,
        "topic_classification": classification,
        "keyword_stats": keyword_stats(raw_rows),
        "raw_crawl": raw_rows,
        "reading_analysis": reading_analysis,
        "archive": archive,
    }


# --------------------------------------------------------------------------- sources


def read_xlsx(path: Path) -> Dict[str, List[List[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - dependency hint
        raise SystemExit("openpyxl is required for --xlsx. Run: pip install -r requirements.txt") from error

    workbook = load_workbook(path, data_only=True, read_only=True)
    tabs: Dict[str, List[List[Any]]] = {}
    for name in TABS + OPTIONAL_TABS:
        if name not in workbook.sheetnames:
            if name in TABS:
                log(f"xlsx is missing tab '{name}', skipping")
            continue
        sheet = workbook[name]
        tabs[name] = [["" if cell is None else cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    return tabs


def google_client(service_account_file: str):
    import gspread
    from google.oauth2.service_account import Credentials

    credentials = Credentials.from_service_account_file(service_account_file, scopes=SCOPES)
    return gspread.authorize(credentials)


def list_weekly_sheets(client, folder_id: str, title_prefix: str) -> List[Dict[str, Any]]:
    files = client.list_spreadsheet_files(folder_id=folder_id)
    sheets = []
    for file in files:
        name = file.get("name", "")
        if title_prefix and not name.startswith(title_prefix):
            continue
        match = DATE_IN_TITLE.search(name)
        sheets.append(
            {
                "id": file.get("id"),
                "title": name,
                "run_date": match.group(1) if match else "",
                "modified_time": file.get("modifiedTime", ""),
                "url": f"https://docs.google.com/spreadsheets/d/{file.get('id')}/edit",
            }
        )
    sheets.sort(key=lambda item: (item["run_date"], item["modified_time"]), reverse=True)
    return sheets


def read_spreadsheet(client, spreadsheet_id: str) -> Dict[str, List[List[Any]]]:
    spreadsheet = client.open_by_key(spreadsheet_id)
    tabs: Dict[str, List[List[Any]]] = {}
    for worksheet in spreadsheet.worksheets():
        if worksheet.title in TABS + OPTIONAL_TABS:
            tabs[worksheet.title] = worksheet.get_all_values()
    missing = [name for name in TABS if name not in tabs]
    if missing:
        log(f"spreadsheet is missing tabs {missing}")
    return tabs


# --------------------------------------------------------------------------- main


def parse_args(argv: Optional[Iterable[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", type=Path, help="Read a local xlsx copy instead of Google Drive")
    parser.add_argument("--spreadsheet-id", help="Read this Google Sheet instead of searching the folder")
    parser.add_argument("--folder-id", help="Drive folder that holds the weekly sheets")
    parser.add_argument("--output", type=Path, default=OUTPUT_PATH, help="Where to write the JSON")
    return parser.parse_args(argv)


def main(argv: Optional[Iterable[str]] = None) -> int:
    load_dotenv(ROOT / ".env")
    args = parse_args(argv)
    title_prefix = os.getenv("CLAUDE_CRAWL_TITLE_PREFIX", DEFAULT_TITLE_PREFIX)

    if args.xlsx:
        if not args.xlsx.exists():
            log(f"xlsx not found: {args.xlsx}")
            return 2
        match = DATE_IN_TITLE.search(args.xlsx.stem)
        source = {
            "kind": "xlsx",
            "path": str(args.xlsx),
            "title": args.xlsx.stem,
            "run_date": match.group(1) if match else "",
            "url": "",
        }
        payload = build_payload(read_xlsx(args.xlsx), source, archive=[])
    else:
        default_service_account = ROOT / "secrets" / "google-service-account.json"
        service_account_file = os.getenv("GOOGLE_SERVICE_ACCOUNT_FILE") or str(default_service_account)
        if not Path(service_account_file).exists() and default_service_account.exists():
            # .env often points at the Docker path (/workspace/secrets/...). Fall back to the repo copy.
            service_account_file = str(default_service_account)
        if not Path(service_account_file).exists():
            log(f"service account file not found: {service_account_file}")
            return 2
        folder_id = args.folder_id or os.getenv("CLAUDE_CRAWL_DRIVE_FOLDER_ID", DEFAULT_FOLDER_ID)
        spreadsheet_id = args.spreadsheet_id or os.getenv("CLAUDE_CRAWL_SPREADSHEET_ID", "")

        try:
            client = google_client(service_account_file)
            archive: List[Dict[str, Any]] = []
            if not spreadsheet_id:
                archive = list_weekly_sheets(client, folder_id, title_prefix)
                if not archive:
                    log(
                        "no weekly sheets visible in the folder. Share the folder with the service account "
                        "(viewer) and make sure the Google Drive API is enabled for the project."
                    )
                    return 3
                spreadsheet_id = archive[0]["id"]
            tabs = read_spreadsheet(client, spreadsheet_id)
            spreadsheet = client.open_by_key(spreadsheet_id)
            match = DATE_IN_TITLE.search(spreadsheet.title)
            source = {
                "kind": "google_sheet",
                "spreadsheet_id": spreadsheet_id,
                "title": spreadsheet.title,
                "run_date": match.group(1) if match else "",
                "url": f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/edit",
                "folder_id": folder_id,
                "folder_url": f"https://drive.google.com/drive/folders/{folder_id}",
            }
        except Exception as error:  # noqa: BLE001 - surface any API failure with a hint
            log(f"Google API error: {error}")
            log("Hints: enable the Google Drive API for the service-account project, and share the Drive folder with the service account email.")
            return 4

        payload = build_payload(tabs, source, archive=archive[1:] if archive else [])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    merged = merge_reading_analysis_into_cache(payload.get("reading_analysis") or [], payload.get("run_date") or "")
    if merged:
        log(f"merged {merged} reading analyses from the routine into {TRANSCRIPT_CACHE_PATH.name}")
    stats = payload["stats"]
    log(
        f"wrote {args.output} (run_date={payload['run_date'] or '?'}, collected={stats['collected']}, "
        f"included={stats['included']}, ideas={stats['ideas']})"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
